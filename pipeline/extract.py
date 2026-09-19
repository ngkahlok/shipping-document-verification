"""Turn a raw SI/BL attachment (txt/pdf/docx/xlsx) into (raw_label, raw_value)
pairs, regardless of format. The caller maps labels to canonical fields via
fields.match_field and normalizes values via normalize.normalized_value.
"""
import io
import re

from fields import match_field, value_plausible


class Unreadable(Exception):
    """Raised when an attachment has no extractable content at all."""


def _pairs_from_lines(lines):
    pairs = []
    for line in lines:
        line = line.rstrip()
        if not line.strip():
            continue
        # Wide gap => two columns rendered side by side (PDF blocks) even
        # with no colon. Try that first since some values (addresses) also
        # legitimately contain colons.
        parts = re.split(r"\s{3,}", line.strip())
        if len(parts) >= 2 and ":" not in parts[0]:
            pairs.append((parts[0], " ".join(parts[1:])))
            continue
        if ":" in line:
            label, value = line.split(":", 1)
            pairs.append((label, value.strip()))
    return pairs


def pairs_from_txt(raw_bytes):
    text = raw_bytes.decode("utf-8", errors="replace")
    if not text.strip():
        raise Unreadable("empty text file")
    return _pairs_from_lines(text.splitlines())


def pairs_from_pdf(raw_bytes):
    import pdfplumber

    try:
        with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
            pairs = []
            any_chars = False
            for page in pdf.pages:
                chars = page.chars
                if chars:
                    any_chars = True
                pairs.extend(_pairs_from_chars(chars))
            if not any_chars:
                raise Unreadable("no text layer (scanned/image-only PDF)")
            return pairs
    except Unreadable:
        raise
    except Exception as e:
        raise Unreadable(f"corrupt/unparseable PDF: {e}")


def _pairs_from_chars(chars, y_tol=2.5):
    """Group characters into visual lines by y-position, then split each line
    into a label/value pair.

    A long label (e.g. "Notify Party/Intermediate Consignee") can visually
    overlap the fixed value column it's paired with -- two separate
    drawString calls whose rendered x-ranges collide. pdfplumber's spatial
    word-clustering garbles that into one interleaved token, but the
    underlying character STREAM order stays clean (every label char is
    emitted before the value's), so we split there instead of by position:
    on the line's colon if it has one (a single drawString call like
    "Container Count: 6 x 40'HC"), else at the first point stream order runs
    backwards in x (a new text object starting left of where the previous
    one had already reached) or, failing that, the widest forward gap.
    """
    rows = {}
    for c in chars:
        if not c.get("text", "").strip() and c.get("text") != " ":
            continue
        key = round(c["top"] / y_tol)
        rows.setdefault(key, []).append(c)

    pairs = []
    for key in sorted(rows):
        cs = rows[key]  # stream order, NOT re-sorted by position
        line = "".join(c["text"] for c in cs).strip()
        if len(cs) < 2 or not line:
            continue
        if ":" in line:
            label, value = line.split(":", 1)
            pairs.append((label, value.strip()))
            continue

        split_i = None
        running_max_x1 = cs[0]["x1"]
        for i in range(1, len(cs)):
            if cs[i]["x0"] < running_max_x1 - 2.0:
                split_i = i
                break
            running_max_x1 = max(running_max_x1, cs[i]["x1"])
        if split_i is None:
            gaps = [(cs[i + 1]["x0"] - cs[i]["x1"], i + 1) for i in range(len(cs) - 1)]
            if not gaps:
                continue
            _, split_i = max(gaps)

        label = "".join(c["text"] for c in cs[:split_i]).strip()
        value = "".join(c["text"] for c in cs[split_i:]).strip()
        if label and value:
            pairs.append((label, value))
    return pairs


def pairs_from_docx(raw_bytes):
    from docx import Document

    try:
        doc = Document(io.BytesIO(raw_bytes))
    except Exception as e:
        raise Unreadable(f"corrupt/unparseable DOCX: {e}")
    pairs = []
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            if len(cells) >= 2:
                pairs.append((cells[0].text, cells[1].text))
    if not pairs:
        # No table -- fall back to paragraph "Label: value" lines.
        pairs = _pairs_from_lines(p.text for p in doc.paragraphs)
    if not pairs:
        raise Unreadable("docx has no readable content")
    return pairs


def pairs_from_xlsx(raw_bytes):
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
    except Exception as e:
        raise Unreadable(f"corrupt/unparseable XLSX: {e}")
    ws = wb.active
    pairs = []
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        vals = [v for v in row if v is not None]
        if len(vals) >= 2:
            pairs.append((str(vals[0]), vals[1]))
    if not pairs:
        raise Unreadable("xlsx has no readable rows")
    return pairs


EXTRACTORS = {
    ".txt": pairs_from_txt,
    ".pdf": pairs_from_pdf,
    ".docx": pairs_from_docx,
    ".xlsx": pairs_from_xlsx,
}


def extract_fields(raw_bytes, filename):
    """Return (fields dict {canonical: raw_value}, raw_pairs list).

    Raises Unreadable if the file is empty or the format can't be parsed at
    all -- callers treat that as the `unreadable` review reason.
    """
    if not raw_bytes:
        raise Unreadable("empty file")
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    extractor = EXTRACTORS.get(ext)
    if extractor is None:
        raise Unreadable(f"unsupported attachment type: {filename}")
    raw_pairs = extractor(raw_bytes)

    fields = {}
    for raw_label, raw_value in raw_pairs:
        field = match_field(raw_label)
        if not field or field in fields:  # first *valid* match wins
            continue
        if not value_plausible(field, raw_value):
            continue
        fields[field] = raw_value
    return fields, raw_pairs
